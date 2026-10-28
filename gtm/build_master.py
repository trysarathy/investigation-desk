"""Build the Master workbook (.xlsx) from CSVs in ./work, ready to upload to Google Drive.

  python -m gtm.build_master --dir work --out Clearline_GTM_Master.xlsx

Upload it to your PERSONAL Google Drive and open it with Google Sheets (File > Save as Google
Sheets), then rename it exactly "Clearline GTM OS — Master" so the scheduled agents find it.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .model import DNC_COLS, LEAD_COLS, LOG_COLS, QUEUE_COLS, STAGES, TOUCH_COLS, read_csv

SETTINGS = [
    ("PAUSE", "FALSE", "TRUE stops all drafting and sending at the next run (the real kill switch)"),
    ("AUTONOMY_LEVEL", "1", "ceiling you allow: 1 = you approve every first email; 2 = Tier 2/3 auto; 3 = all auto. The guard lowers it automatically on any downgrade trigger"),
    ("DAILY_CAP", "30", "hard cap on first emails per day"),
    ("CURRENT_SEND_CAP", "5", "today's real cap while sending from personal Gmail; raise to 15-30 only with your own domain"),
    ("SIMILAR_30D_CAP", "1000", "max similar emails in any 30 days"),
    ("SAME_COMPANY_DAYS", "7", "never two people at one company within this many days"),
    ("SENDER", "aashmiin@gmail.com", "change to your domain inbox when you have one"),
    ("BOOKING_LINK", "", "paste your Google Calendar booking page; until then 'yes' replies are drafted for you"),
    ("POSTAL_ADDRESS", "", "required before emailing US contacts (CAN-SPAM)"),
    ("OWNER_TZ", "Asia/Singapore", ""),
    ("FOLLOWUP_DAYS", "4,11", "follow-ups go in the same thread"),
    ("MASTER_REPO", "https://github.com/trysarathy/clearline-gtm", "guard script lives here"),
    ("SEND_MODE", "gmail_personal", "gmail_personal = send from aashmiin@gmail.com via the Gmail connector (cap above). domain = separate domain inboxes through the sending tool, with the ramp below"),
    ("INBOXES", "", "domain inboxes, comma-separated, e.g. aashmiin@getclearline.com,a.kaur@getclearline.com"),
    ("PER_INBOX_CAP", "20", "max emails per inbox per day (first + follow-ups)"),
    ("RAMP_START", "", "YYYY-MM-DD: first day of warm-up week (week 1 = warm-up only)"),
    ("RAMP_CAPS", "0,20,35,50", "emails/day for week 1, 2, 3, 4+ (first + follow-ups combined)"),
    ("DOMAIN_READY", "FALSE", "TRUE only after SPF, DKIM and DMARC pass"),
    ("UNSUBSCRIBE_READY", "FALSE", "TRUE only after a one-click unsubscribe link is in every email"),
    ("PLACEMENT_OK", "FALSE", "TRUE when the latest Gmail + Outlook inbox-placement test landed in the inbox; needed for each step up"),
    ("BULK_CONFIRMED", "FALSE", "TRUE only after a professional confirms Spam Control Act bulk duties; sending stops at 1,000 in 30 days until then"),
    ("THROTTLE_CAP", "20", "auto-throttle level: bounce >3%, any complaint, or reply <1% over 100 sends"),
]

DASHBOARD = [
    ("Leads in Master", "=COUNTA(Leads!A2:A5000)", "rows in Leads"),
    ("Tier 1 leads", '=COUNTIF(Leads!B2:B5000,"1")', ""),
    ("Leads added (last 7 days)", '=COUNTIFS(Log!C2:C5000,"lead added",Log!A2:A5000,">="&TEXT(TODAY()-7,"yyyy-mm-dd"))', "from Log"),
    ("First emails sent", '=COUNTIFS(Touches!G2:G5000,"D0",Touches!I2:I5000,"sent")', "Touches: step D0, status sent"),
    ("Bounce %", '=IFERROR(COUNTIF(Touches!I2:I5000,"bounced")/B5,0)', "bounced ÷ first emails"),
    ("Reply %", '=IFERROR(COUNTIF(Touches!I2:I5000,"replied")/B5,0)', "replied ÷ first emails"),
    ("Positive reply %", '=IFERROR(COUNTIFS(Touches!I2:I5000,"replied",Touches!L2:L5000,"positive*")/B5,0)', "replies whose notes start with 'positive'"),
    ("Calls booked", '=COUNTIF(Calls!E2:E5000,"booked")+COUNTIF(Calls!E2:E5000,"held")', ""),
    ("Calls held", '=COUNTIF(Calls!E2:E5000,"held")', ""),
    ("Strong problems", '=COUNTIF(Calls!F2:F5000,"Strong")', "problem_score on calls"),
    ("Design partners", '=COUNTIF(Leads!D2:D5000,"Design partner")', ""),
    ("Paying", '=COUNTIF(Leads!D2:D5000,"Paying")', ""),
    ("Emails sent today (all steps)", '=COUNTIFS(Touches!A2:A5000,TEXT(TODAY(),"yyyy-mm-dd"),Touches!I2:I5000,"sent")', "first + follow-ups"),
    ("Emails sent, last 30 days", '=COUNTIFS(Touches!A2:A5000,">="&TEXT(TODAY()-29,"yyyy-mm-dd"),Touches!I2:I5000,"sent")', "Spam Control Act bulk line = 1,000"),
    ("Waiting for your approval", "=COUNTIFS('Approval Queue'!O2:O5000,\"queued\")", ""),
    ("PAUSE", "=Settings!B2", ""),
    ("Autonomy level allowed", "=Settings!B3", ""),
    ("Decision rule (after 10 calls)", "≥4 Strong → pilots · 2-3 → narrow segment · ≤1 → rethink", ""),
]

H, F = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2A44")


def _tab(wb, name, cols, rows=(), widths=None, wrap=()):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = H, F
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    ws.freeze_panes = "B2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 16)
        if c in wrap:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def build(work: Path, out: Path) -> None:
    leads = read_csv(work / "leads.csv")
    leads.sort(key=lambda x: (x.get("tier", "9"), -int(x.get("score") or 0)))
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in [
        ["Clearline GTM OS — Master (keep in your PERSONAL Drive, never the Hatch account)"], [],
        ["Approve emails", "Approval Queue tab: set APPROVE to yes (or edit + write edited_body). To approve all, fill the column with yes."],
        ["Kill switch", "Settings › PAUSE = TRUE stops all drafting and sending at the next scheduled run."],
        ["Autonomy", "Settings › AUTONOMY_LEVEL is the most you allow; the guard drops it to 1 automatically on any bad signal."],
        ["LinkedIn", "Never automated. The morning briefing lists people for you to contact by hand."],
        ["Stages", " → ".join(STAGES)],
        ["Rule", "Every first line needs a source URL. No source = no send."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 110
    _tab(wb, "Settings", ["key", "value", "meaning"], [dict(key=k, value=v, meaning=m) for k, v, m in SETTINGS],
         {"key": 20, "value": 40, "meaning": 100})
    lt = _tab(wb, "Leads", LEAD_COLS, leads, {"company": 30, "first_line": 60, "check_before_sending": 50, "email": 30,
                                               "evidence_url": 40, "first_line_source": 40, "segment": 30})
    dv = DataValidation(type="list", formula1='"' + ",".join(STAGES) + '"', allow_blank=True)
    lt.add_data_validation(dv)
    dv.add("D2:D5000")
    q = _tab(wb, "Approval Queue", QUEUE_COLS, read_csv(work / "queue.csv"),
             {"body": 90, "edited_body": 60, "guard_reasons": 40, "subject": 30, "first_line_source": 40}, wrap=("body",))
    dv2 = DataValidation(type="list", formula1='"yes,no,edit"', allow_blank=True)
    q.add_data_validation(dv2)
    dv2.add("M2:M2000")
    _tab(wb, "Touches", TOUCH_COLS, read_csv(work / "touches.csv"), {"notes": 40, "subject": 30, "email": 30})
    _tab(wb, "Calls", ["date", "lead_id", "company", "person", "status", "problem_score", "pain_quotes",
                       "current_workaround", "volume_per_month", "next_step", "notes_link"], read_csv(work / "calls.csv"),
         {"pain_quotes": 60})
    _tab(wb, "Opportunities", ["lead_id", "company", "stage", "pilot_scope", "started", "next_step", "next_step_date", "notes"])
    _tab(wb, "Do-not-contact", DNC_COLS, read_csv(work / "dnc.csv"), {"value": 34, "reason": 40})
    _tab(wb, "Spot checks", ["date", "lead_id", "email_subject", "checked_fact", "source_url", "factual_error", "notes"],
         read_csv(work / "spot_checks.csv"), {"checked_fact": 50, "source_url": 40})
    _tab(wb, "Log", LOG_COLS, read_csv(work / "log.csv"), {"details": 90, "timestamp": 20})
    d = wb.create_sheet("Dashboard")
    d.append(["Metric", "Value", "How it's counted"])
    for c in d[1]:
        c.font, c.fill = H, F
    for r in DASHBOARD:
        d.append(list(r))
    for cell in ("B6", "B7", "B8"):
        d[cell].number_format = "0.0%"
    d.column_dimensions["A"].width, d.column_dimensions["B"].width, d.column_dimensions["C"].width = 32, 24, 60
    wb.save(out)


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--dir", default="work")
    p.add_argument("--out", default="Clearline_GTM_Master.xlsx")
    a = p.parse_args()
    build(Path(a.dir), Path(a.out))
    print(a.out)
