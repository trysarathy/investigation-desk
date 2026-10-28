"""Shared look of the Clearline decks: dark, Arial, ink / white / periwinkle (red only for warnings),
plus the slide helpers. Import with `from deck_style import *`."""
import sys, pathlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

A = pathlib.Path(__file__).resolve().parent.parent / "assets"
FONT = "Arial"
INK, CARD, LINE = RGBColor(0x0B, 0x0E, 0x14), RGBColor(0x15, 0x1A, 0x24), RGBColor(0x2A, 0x31, 0x40)
WHITE, MUTED, FAINT = RGBColor(0xF2, 0xF4, 0xF7), RGBColor(0xA3, 0xAA, 0xB6), RGBColor(0x76, 0x7E, 0x8C)
ACC, ACC2, RED = RGBColor(0x8E, 0xA2, 0xFF), RGBColor(0xB9, 0xC6, 0xFF), RGBColor(0xFF, 0x8A, 0x7F)
W, H, M = 13.333, 7.5, 0.6

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]

def new(notes):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
    s.notes_slide.notes_text_frame.text = notes.strip()
    return s

def text(s, x, y, w, h, runs, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, name=None):
    """runs: a string, or a list of paragraphs; each paragraph a string or a list of (text, {size,color,bold}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font
            f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color); f.italic = o.get("italic", False)
    return tb

def title(s, t, sub=None):
    text(s, M, 0.5, W - 2 * M, 0.9, t, size=36, bold=True, name="Title")
    if sub: text(s, M, 1.38, W - 2 * M, 0.5, sub, size=20, color=MUTED, name="Subtitle")

def footer(s, t):
    text(s, M, H - 0.5, W - 2 * M, 0.3, t, size=10, color=FAINT, name="Sources")

def box(s, x, y, w, h, fill=CARD, line=LINE, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius; sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh

def pill(s, x, y, t, color=ACC, size=12, w=None):
    w = w or (0.18 + 0.085 * len(t) * size / 12)
    b = box(s, x, y, w, 0.34 * size / 12, fill=INK, line=color, radius=0.5)
    text(s, x, y, w, 0.34 * size / 12, t, size=size, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w

def img(s, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)

def fit(path, maxw, maxh):
    im = Image.open(path); r = min(maxw / im.width, maxh / im.height); return im.width * r, im.height * r

TAG = {"A": ("ASSUMPTION", ACC2), "G": ("GUESS", ACC2), "V": ("VERIFIED", MUTED)}
def tag(kind, n=""):
    t, c = TAG[kind]; return (f"  [{t}{(' ' + n) if n else ''}]", {"size": 12, "color": c, "bold": False})

